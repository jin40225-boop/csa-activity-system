"""產生暫用的「講師領據」Word 範本（佔位字每個都在同一段文字裡，系統填值不會漏）。

正式版：拿貴中心現行領據 .docx，把要填的地方改成 {{佔位字}}，存成 範本\\講師領據.docx 覆蓋即可。
用法：python tools/make_receipt_template.py  → 輸出 範本/講師領據.docx
"""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

out = Path(__file__).resolve().parents[1] / '範本' / '講師領據.docx'
doc = Document()
style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(13)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '標楷體')


def para(text, size=13, bold=False, align=None, color=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    if align:
        p.alignment = align
    return p


para('{{方案名稱}}', 14, align=WD_ALIGN_PARAGRAPH.CENTER)
para('講師鐘點費領據', 20, True, WD_ALIGN_PARAGRAPH.CENTER)
para('（暫用格式：請以貴中心現行領據範本取代本檔，保留兩層大括號的佔位字即可）', 10, align=WD_ALIGN_PARAGRAPH.CENTER, color='888888')

rows = [
    ('課程名稱', '{{課程名稱}}'),
    ('辦理日期', '{{辦理日期}}　{{辦理時間}}'),
    ('辦理地點', '{{辦理地點}}'),
    ('講師姓名', '{{講師姓名}}'),
    ('授課時數', '{{授課時數}} 小時'),
    ('鐘點費', '每小時 {{鐘點費單價}} 元 × {{授課時數}} 小時 ＝ {{鐘點費}} 元'),
    ('交通費', '{{交通費}}'),
    ('合計', '新臺幣 {{合計大寫}}（{{合計}} 元）'),
    ('身分證統一編號', ''),
    ('戶籍地址', ''),
    ('領款人簽章', ''),
]
t = doc.add_table(rows=len(rows), cols=2)
t.style = 'Table Grid'
t.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, (k, v) in enumerate(rows):
    t.cell(i, 0).text = k
    t.cell(i, 1).text = v
    t.cell(i, 0).paragraphs[0].runs[0].bold = True
    for c in (0, 1):
        t.cell(i, c).width = Pt(110 if c == 0 else 330)

para('')
para('上開款項業已如數收訖無誤。')
para('中華民國 {{民國年}} 年 {{月}} 月 {{日}} 日', align=WD_ALIGN_PARAGRAPH.RIGHT)
sig = doc.add_table(rows=2, cols=3)
sig.style = 'Table Grid'
for i, k in enumerate(['承辦人 {{承辦}}', '督導', '會計']):
    sig.cell(0, i).text = k
out.parent.mkdir(exist_ok=True)
doc.save(out)
print(out)
